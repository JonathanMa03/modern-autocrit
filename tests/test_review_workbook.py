import json
from pathlib import Path

from openpyxl import load_workbook

from backend.batch_manifest import manifest_state_path
from backend.review_workbook import _bool_or_none, export_review_workbook


def test_export_review_workbook_contains_completed_rows(tmp_path: Path) -> None:
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps({"cohort_id": "test-review", "trials": [{"trial_id": "NCT00000001"}]}),
        encoding="utf-8",
    )
    runtime = tmp_path / "runtime" / "criteria"
    jobs = runtime / "jobs"
    jobs.mkdir(parents=True)
    job_id = "job-one"
    row = {
        "criterion_id": "criterion-one", "row_id": "row-one", "source_id": "source-one",
        "trial_id": "NCT00000001", "criterion_type": "inclusion", "domain": "person",
        "entity": "age", "comparator": ">=", "value": "18", "unit": "years",
        "source": "Age at least 18 years.", "surrounding_context": "Inclusion: Age at least 18 years.",
        "human_review_status": "pending", "validation_issues": [], "readiness_blockers": ["human_review_required"],
    }
    (jobs / f"{job_id}.json").write_text(
        json.dumps({"job_id": job_id, "trial_id": "NCT00000001", "status": "completed", "rows": [row]}),
        encoding="utf-8",
    )
    state_path = manifest_state_path(manifest, runtime)
    state_path.parent.mkdir(parents=True)
    state_path.write_text(
        json.dumps({"jobs": {"NCT00000001": {"job_id": job_id}}}), encoding="utf-8"
    )
    output = tmp_path / "review.xlsx"

    result = export_review_workbook(manifest, tmp_path, output)
    workbook = load_workbook(output)
    sheet = workbook["Human Review"]

    assert result["completed_jobs"] == 1
    assert result["rows"] == 1
    assert sheet.max_row == 2
    headers = [cell.value for cell in sheet[1]]
    assert workbook.sheetnames[:2] == ["Validation Instructions", "Human Review"]
    assert sheet.cell(2, 1).value == "Age at least 18 years."
    assert sheet.cell(2, 2).value is None
    assert sheet.cell(2, headers.index("trial_id") + 1).value == "NCT00000001"
    assert sheet.cell(2, headers.index("review_status") + 1).value == "pending"
    assert sheet.freeze_panes == "C2"
    guide = workbook["Validation Instructions"]
    labels = {guide.cell(row, 1).value: guide.cell(row, 2).value for row in range(1, guide.max_row + 1)}
    assert labels["Completed trials"] == 1
    assert labels["Extracted criteria"] == 1


def test_review_boolean_parser_handles_excel_and_text_values() -> None:
    assert _bool_or_none(True) is True
    assert _bool_or_none("FALSE") is False
    assert _bool_or_none(1) is True
    assert _bool_or_none(None) is None
