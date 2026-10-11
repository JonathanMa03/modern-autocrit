import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.eav_pipeline import canonical_term, deterministic_criterion, deterministic_text_constraints, normalize_rows, segment_criteria
from backend.criterion_validation import readiness, validate_criterion
from backend.services.eav_workflow import EavWorkflowService, Job


def test_numerical_criterion_separates_fields_and_repetition():
    row = deterministic_criterion(
        source="Blood pressure >145/95 on three occasions",
        criterion_type="exclusion", entity="blood pressure >145/95 on three occasions",
    )
    assert row == {
        "criterion_type": "exclusion", "domain": "measurement",
        "entity": "Blood Pressure", "comparator": ">", "value": "145/95",
        "unit": None, "value_kind": "numerical", "negated": False,
        "lower_bound": None, "lower_inclusive": None,
        "upper_bound": None, "upper_inclusive": None, "interval": None,
        "temporal": None, "repetition": "on three occasions", "qualifier": None,
        "logical_operator": "standalone",
        "relations": ["HAS_VALUE", "HAS_MULTIPLIER"],
        "source": "Blood pressure >145/95 on three occasions",
    }


def test_unit_is_separate_from_value():
    row = deterministic_criterion(
        source=r"dRVVT \>37 sec", criterion_type="exclusion", entity="dRVVT >37 sec"
    )
    assert (row["entity"], row["comparator"], row["value"], row["unit"]) == (
        "dRVVT", ">", "37", "sec"
    )


def test_categorical_negation_is_explicit_and_informative():
    row = deterministic_criterion(
        source="no history of menopause", criterion_type="inclusion", entity="menopause"
    )
    assert row["entity"] == "menopause"
    assert row["comparator"] is None
    assert row["value"] == "absent"
    assert row["negated"] is True


def test_model_categorical_value_is_preserved_when_no_comparator():
    row = deterministic_criterion(
        source="Both male and female participants are eligible",
        criterion_type="inclusion", entity="sex", value="male and female",
    )
    assert row["entity"] == "Sex"
    assert row["value"] == "male and female"


def test_embedded_not_due_to_does_not_negate_entity():
    row = deterministic_criterion(
        source="Diabetes mellitus (NOT due to steroids) with vascular disease",
        criterion_type="exclusion", entity="diabetes mellitus with vascular disease",
    )
    assert row["value"] == "present"
    assert row["negated"] is False


def test_temporal_construct_is_separate():
    row = deterministic_criterion(
        source="Myocardial infarction within the past 60 days",
        criterion_type="exclusion", entity="myocardial infarction",
    )
    assert row["temporal"] == "within the past 60 days"
    assert row["entity"] == "myocardial infarction"


def test_segmentation_has_stable_source_ids_and_splits_semicolons():
    rows = segment_criteria("Exclusion Criteria:\nGPL >40; MPL >40; APL >50; dRVVT >37 sec")
    assert [row["source"] for row in rows] == ["GPL >40", "MPL >40", "APL >50", "dRVVT >37 sec"]
    assert all(row["criterion_type"] == "exclusion" for row in rows)
    assert len({row["source_id"] for row in rows}) == 1
    assert rows[0]["source_id"].startswith("statement-")
    assert rows == segment_criteria("Exclusion Criteria:\nGPL >40; MPL >40; APL >50; dRVVT >37 sec")


def test_segmentation_accepts_qualified_registry_headings():
    rows = segment_criteria(
        "General Inclusion Criteria:\n* ECOG performance status of 0 or 1\n"
        "Inclusion Criteria for Participants with Stage III NSCLC:\n* Unresectable Stage III NSCLC\n"
        "Exclusion Criteria Trial and Registry:\n* Prior myocardial infarction"
    )
    assert [(row["criterion_type"], row["source"]) for row in rows] == [
        ("inclusion", "ECOG performance status of 0 or 1"),
        ("inclusion", "Unresectable Stage III NSCLC"),
        ("exclusion", "Prior myocardial infarction"),
    ]


def test_dense_legacy_registry_fields_are_separate_parent_statements():
    rows = segment_criteria(
        "Inclusion Criteria:\n* PATIENT CHARACTERISTICS: Age: 55 and under "
        "Performance status: Karnofsky 70-100% Life expectancy: Not specified "
        "Hepatic: Bilirubin less than 2.5 mg/dL Renal: Creatinine within normal range"
    )
    assert [row["source"].split(":", 1)[0].casefold() for row in rows] == [
        "age", "performance status", "life expectancy", "hepatic", "renal"
    ]
    assert len({row["parent_statement_id"] for row in rows}) == 5


def test_multiple_entities_share_source_and_retain_boolean_logic():
    source = "deep vein thrombosis or pulmonary embolus"
    rows = normalize_rows({"rows": [
        _model_row(source, "DVT", "condition"),
        _model_row(source, "pulmonary embolus", "condition"),
    ]}, [{"criterion_type": "exclusion", "source": source, "source_id": "source-0007"}])
    assert len(rows) == 2
    assert {row["logical_operator"] for row in rows} == {"or"}
    assert all("OR" in row["relations"] for row in rows)
    assert {row["source_id"] for row in rows} == {"source-0007"}


def test_entity_relative_comparator_handles_two_measurements_in_one_source():
    source = "GPL>50 and blood pressure >145/95"
    rows = normalize_rows({"rows": [
        _model_row(source, "GPL", "measurement", comparator=">", value="50", logic="and"),
        _model_row(source, "blood pressure", "measurement", comparator=">", value="145/95", logic="and"),
    ]}, [{"criterion_type": "exclusion", "source": source, "source_id": "source-0001"}])
    assert [(row["entity"], row["value"]) for row in rows] == [("GPL", "50"), ("Blood Pressure", "145/95")]


def test_compound_measurements_are_grounded_to_nearest_labels():
    source = "Uncontrolled hypertension (systolic BP >180 mmHg, diastolic BP >110 mmHg)"
    rows = normalize_rows({"rows": [
        _model_row(source, "uncontrolled hypertension", "measurement", comparator=">", value="180"),
        _model_row(source, "uncontrolled hypertension", "measurement", comparator=">", value="110"),
    ]}, [{"criterion_type": "exclusion", "source": source, "source_id": "source-bp"}])
    assert [(row["entity"], row["value"]) for row in rows] == [
        ("systolic Blood Pressure", "180"), ("diastolic Blood Pressure", "110")
    ]
    assert not any(
        issue["code"] == "DUPLICATE_ATOMIC_CRITERION"
        for row in rows for issue in row["validation_issues"]
    )


def test_duplicate_atomic_rows_are_collapsed_deterministically():
    source = "Platelets >=100x10^9/L"
    rows = normalize_rows({"rows": [
        _model_row(source, "platelets", "measurement", comparator=">=", value="100"),
        _model_row(source, "platelets", "condition", comparator=">=", value="100"),
    ]}, [{"criterion_type": "inclusion", "source": source, "source_id": "source-platelets"}])
    assert len(rows) == 1
    assert rows[0]["domain"] == "measurement"


def test_literal_comparator_overrides_contradictory_model_bounds():
    row = deterministic_criterion(
        source="Serum creatinine >3 mg/dL", criterion_type="exclusion",
        entity="serum creatinine", comparator=">", value=3, unit="mg/dL",
        lower_bound=None, upper_bound=3, upper_inclusive=False,
    )
    assert row["comparator"] == ">"
    assert row["lower_bound"] == 3 and row["lower_inclusive"] is False
    assert row["upper_bound"] is None and row["upper_inclusive"] is None
    assert row["interval"] == "(3, +inf)"


def test_repeated_values_use_entity_anchor_for_karnofsky_and_lansky():
    source = (
        "Karnofsky score >50% for patients >16 years of age and "
        "Lansky >50% for patients <=16 years of age"
    )
    rows = normalize_rows({"rows": [
        _model_row(source, "Karnofsky score", "measurement", comparator=">", value="50", logic="and"),
        _model_row(source, "Lansky score", "measurement", comparator=">", value="50", logic="and"),
    ]}, [{"criterion_type": "inclusion", "source": source, "source_id": "source-performance"}])
    assert [(row["entity"], row["value"]) for row in rows] == [
        ("Karnofsky Performance Status", "50"), ("lansky score", "50")
    ]
    assert not any(
        issue["code"] == "DUPLICATE_ATOMIC_CRITERION"
        for row in rows for issue in row["validation_issues"]
    )


def test_grade_comparison_is_parsed_without_model_help():
    row = deterministic_criterion(
        source="Peripheral sensory neuropathy > grade 1",
        criterion_type="exclusion", entity="peripheral sensory neuropathy",
    )
    assert (row["entity"], row["comparator"], row["value"], row["unit"]) == (
        "peripheral sensory neuropathy", ">", "1", "grade"
    )


@pytest.mark.parametrize("source, expected", [
    ("Disease progressed during/following last prior treatment", "during/following last prior treatment"),
    ("Discussed with the patient before registration in the trial", "before registration in the trial"),
    ("Atrial fibrillation at any time, but before the screening visit", "before the screening visit"),
])
def test_explicit_temporal_phrases_are_deterministic(source, expected):
    row = deterministic_criterion(
        source=source, criterion_type="inclusion", entity="clinical event",
    )
    assert expected in row["temporal"]


def test_written_comparators_are_extracted_without_model_help():
    statins = deterministic_text_constraints(
        "atorvastatin greater than or equal to 40 mg or rosuvastatin greater than or equal to 20 mg",
        "inclusion",
    )
    assert [(row["entity"], row["comparator"], row["value"], row["unit"]) for row in statins] == [
        ("atorvastatin", ">=", 40, "mg"), ("rosuvastatin", ">=", 20, "mg")
    ]
    assert {row["logical_operator"] for row in statins} == {"or"}
    count = deterministic_text_constraints(
        "Patients on four or more antihypertensive drugs at Visit 1.", "exclusion"
    )
    assert [(row["entity"], row["comparator"], row["value"]) for row in count] == [
        ("antihypertensive drugs", ">=", 4)
    ]
    cycles = deterministic_text_constraints(
        "No more than 2 cycles of prior 5-FU chemotherapy", "inclusion"
    )
    assert cycles[0]["entity"] == "prior 5-FU chemotherapy cycles"
    normalized = deterministic_criterion(**{
        key: value for key, value in cycles[0].items()
        if key in {"source", "criterion_type", "entity", "domain", "comparator", "value", "unit", "negated", "qualifier", "temporal", "repetition", "logical_operator", "lower_bound", "lower_inclusive", "upper_bound", "upper_inclusive"}
    })
    assert normalized["negated"] is False
    assert normalized["interval"] == "(-inf, 2]"
    comma_number = deterministic_text_constraints(
        "Absolute blast count greater than 50,000/mcl", "inclusion"
    )
    assert [(row["entity"], row["value"], row["unit"]) for row in comma_number] == [
        ("Absolute blast count", 50000, "/mcl")
    ]
    escaped_comma = deterministic_criterion(
        source=r"Thrombocytopenia defined as platelet count \<50,000 platelets/μL.",
        criterion_type="exclusion", entity="platelet count", domain="measurement",
    )
    assert (escaped_comma["comparator"], escaped_comma["value"], escaped_comma["unit"]) == (
        "<", "50000", "platelets/μL"
    )
    assert deterministic_text_constraints(
        "No more than 2 cycles of prior 5-FU chemotherapy required. No prior radiation.",
        "inclusion",
    )[0]["entity"] == "prior 5-FU chemotherapy cycles"
    assert deterministic_text_constraints(
        "The expected failure rate is less than 1% per year.", "exclusion"
    )[0]["entity"] == "failure rate"
    assert deterministic_text_constraints(
        "The above medications must be stable for at least 3 months.", "inclusion"
    )[0]["entity"] == "medications stability duration"
    assert deterministic_text_constraints(
        "At least 1 of the following conditions applies: diabetes or smoking.", "inclusion"
    )[0]["entity"] == "following conditions"


def test_deterministic_sibling_satisfies_parent_comparison_validation():
    source = "Severe disease expected to reduce life expectancy to less than 5 years"
    rows = normalize_rows(
        {"rows": [_model_row(source, "severe disease", "condition")]},
        [{"criterion_type": "exclusion", "source": source, "source_id": "source-life"}],
    )
    assert any(row["comparator"] == "<" and row["value"] == "5" for row in rows)
    assert not any(
        issue["code"] == "COMPARISON_NOT_STRUCTURED"
        for row in rows for issue in row["validation_issues"]
    )


def test_age_range_phrasings_normalize_to_same_closed_interval():
    sources = ["aged 18 to 65", "between 18 and 65 years", "18–65 years of age", "age ≥18 and ≤65 years"]
    rows = [deterministic_criterion(source=source, criterion_type="inclusion", entity="Age") for source in sources]
    for row in rows:
        assert row["entity"] == "Age"
        assert row["domain"] == "person"
        assert row["value_kind"] == "interval"
        assert (row["lower_bound"], row["upper_bound"]) == (18, 65)
        assert row["lower_inclusive"] is True and row["upper_inclusive"] is True
        assert row["interval"] == "[18, 65]"
        assert "HAS_VALUE" in row["relations"]


def test_age_mixed_inclusive_range_recovers_entity_and_bounds():
    row = deterministic_criterion(
        source=r"\>= 3 years of age and \<30 years old at study entry",
        criterion_type="inclusion", entity="",
    )
    assert row["entity"] == "Age"
    assert row["unit"] == "years"
    assert (row["lower_bound"], row["upper_bound"]) == (3, 30)
    assert row["lower_inclusive"] is True
    assert row["upper_inclusive"] is False


def test_mixed_month_year_interval_is_normalized_before_order_validation():
    row = deterministic_criterion(
        source="duration of >=6 months and <=3 years",
        criterion_type="inclusion", entity="duration",
    )
    assert (row["lower_bound"], row["upper_bound"], row["unit"]) == (6, 36, "months")
    assert row["interval"] == "[6, 36]"


def test_registry_slash_equal_operator_and_context_free_blood_pressure():
    threshold = deterministic_criterion(
        source=r"Leukocytes \>/= 3,000/microgram", criterion_type="inclusion",
        entity="leukocytes",
    )
    assert (threshold["comparator"], threshold["value"]) == (">=", "3000")
    pressure = deterministic_criterion(
        source=r"(\< 140/90 mmHg)", criterion_type="inclusion", entity="",
    )
    assert pressure["entity"] == "Blood Pressure"


def test_model_negation_cannot_reverse_plain_numerical_threshold():
    row = deterministic_criterion(
        source="eGFR <=15 mL/min/1.73 m2", criterion_type="exclusion",
        entity="eGFR", negated=True,
    )
    assert row["negated"] is False
    assert row["value"] == "15"
    assert row["unit"] == "mL/min/1.73 m2"


def test_temporal_range_is_not_misread_as_entity_value():
    row = deterministic_criterion(
        source="Ischemic stroke within 2-12 weeks after randomization",
        criterion_type="inclusion", entity="ischemic stroke",
    )
    assert row["value"] == "present"
    assert row["lower_bound"] is None and row["upper_bound"] is None
    assert row["temporal"] == "within 2-12 weeks after randomization"


def test_temporal_unit_does_not_trigger_missing_measurement_unit():
    row = {
        "criterion_type": "inclusion", "entity": "neutralizing antibody titer",
        "source": "titer <1:2 or equivocal within 90 days of screening",
        "comparator": "<", "value": "1", "unit": None, "negated": False,
        "temporal": "within 90 days of screening", "repetition": None,
        "lower_bound": None, "lower_inclusive": None, "upper_bound": 1,
        "upper_inclusive": False, "logical_operator": "standalone",
        "parent_atom_count": 1,
    }
    assert "UNIT_UNCAPTURED" not in {item["code"] for item in validate_criterion(row)}


def test_protocol_hash_is_carried_into_normalized_rows():
    row = normalize_rows(
        {"rows": [_model_row("Age >=18 years", "age", "person", comparator=">=", value="18")]},
        protocol_id="NCT1", source_version="v1", protocol_hash="sha256:test",
        source_retrieved_at="2026-01-01T00:00:00+00:00", trial_id="NCT00000001",
    )[0]
    assert row["protocol_hash"] == "sha256:test"
    assert row["trial_id"] == "NCT00000001"
    assert "PROVENANCE_INCOMPLETE" not in row["readiness_blockers"]


def test_configured_aliases_normalize_abbreviations_and_full_names():
    assert canonical_term("kidney impairment") == "renal impairment"
    assert canonical_term("acquired immunodeficiency syndrome") == "AIDS"
    assert canonical_term("Unequivocal diagnosis of systemic lupus erythematosus") == "unequivocal diagnosis of SLE"
    assert canonical_term("IgG phospholipid unit") == "GPL"
    assert canonical_term("NSCLC") == "Non-Small Cell Lung Cancer"
    assert canonical_term("ECOG") == "ECOG Performance Status"
    assert canonical_term("ALL") == "Acute Lymphoblastic Leukemia"
    assert canonical_term("all patients") == "all patients"


def test_legacy_library_rows_are_upgraded_in_memory(tmp_path: Path):
    base = tmp_path / "base.json"
    base.write_text(json.dumps({"immutable": True, "entries": []}), encoding="utf-8")
    legacy = tmp_path / "eavLibrary_20200101_000000_000000.json"
    legacy.write_text(json.dumps({"entries": [{
        "row_id": "old", "trial_id": "NCT1", "criterion_type": "exclusion",
        "entity": "age", "attribute": ">", "value": "65", "source": "Age >65",
    }]}), encoding="utf-8")
    service = EavWorkflowService(object(), tmp_path / "runtime", base)
    row = service.library()["entries"][0]
    assert (row["domain"], row["comparator"], row["value_kind"]) == ("person", ">", "numerical")


def test_review_creates_versioned_structured_json_and_csv(tmp_path: Path):
    base = tmp_path / "base.json"
    base.write_text(json.dumps({"immutable": True, "entries": []}), encoding="utf-8")
    service = EavWorkflowService(object(), tmp_path / "runtime", base)
    row = normalize_rows({"rows": [_model_row("GPL>50", "GPL", "measurement", comparator=">", value="50")]})[0]
    job = Job(job_id="job", trial_id="NCT1", input_mode="manual", status="completed", rows=[row])
    job.protocol_id = "NCT1"
    job.source_version = "v1"
    job.source_retrieved_at = "2026-01-01T00:00:00+00:00"
    job.protocol_hash = "sha256:test"
    service._save(job)
    result = service.apply_review("job", [{**row, "review_status": "accepted"}])
    assert result["json"].startswith("criteriaLibrary_")
    saved = json.loads((service.libraries / result["json"]).read_text())
    assert saved["schema_version"] == "modern-autocrit.structured-criteria.v4"
    assert saved["entries"][0]["domain"] == "measurement"
    assert saved["entries"][0]["interval"] == "(50, +inf)"
    assert saved["entries"][0]["human_review_status"] == "accepted"
    assert saved["entries"][0]["export_ready"] is True
    assert (base.parent / result["csv"]).is_file()


def test_validation_blocks_uncaptured_semantics_and_requires_review():
    row = {
        "protocol_id": "NCT1", "source_version": "v1",
        "criterion_type": "exclusion", "entity": "blood pressure",
        "source": "Blood pressure >145 mmHg twice within three months",
        "comparator": None, "value": "present", "unit": None,
        "negated": False, "temporal": None, "repetition": None,
        "lower_bound": None, "lower_inclusive": None,
        "upper_bound": None, "upper_inclusive": None,
        "logical_operator": "standalone", "parent_atom_count": 1,
        "human_review_status": "pending",
    }
    row["validation_issues"] = validate_criterion(row)
    codes = {item["code"] for item in row["validation_issues"]}
    assert {"COMPARISON_NOT_STRUCTURED", "TEMPORAL_UNCAPTURED"} <= codes
    ready, blockers = readiness(row)
    assert ready is False
    assert "HUMAN_REVIEW_REQUIRED" in blockers


def test_normalized_rows_have_stable_provenance_and_parent_scope():
    source = "pregnant or breastfeeding"
    items = segment_criteria(
        f"Exclusion Criteria:\n- {source}", protocol_id="NCT1", source_version="2026-01-01"
    )
    rows = normalize_rows({"rows": [
        _model_row(source, "pregnancy", "condition", logic="or"),
        _model_row(source, "breastfeeding", "observation", logic="or"),
    ]}, items, protocol_id="NCT1", source_version="2026-01-01")
    assert len({row["parent_statement_id"] for row in rows}) == 1
    assert [row["atom_index"] for row in rows] == [1, 2]
    assert all(row["parent_atom_count"] == 2 for row in rows)
    assert all(row["export_ready"] is False for row in rows)


def test_library_excel_export_contains_base_and_reviewed_rows(tmp_path: Path):
    base = tmp_path / "base.json"
    base.write_text(json.dumps({"immutable": True, "entries": [{
        "row_id": "base", "trial_id": "BASE", "criterion_type": "inclusion",
        "domain": "person", "entity": "Age", "comparator": ">=", "value": "18",
        "unit": "years", "source": "Age >=18", "relations": ["HAS_VALUE"],
    }]}), encoding="utf-8")
    snapshot = tmp_path / "criteriaLibrary_20200101_000000_000000.json"
    snapshot.write_text(json.dumps({"entries": [{
        "row_id": "reviewed", "trial_id": "NCT1", "criterion_type": "exclusion",
        "domain": "condition", "entity": "diabetes", "comparator": None,
        "value": "present", "unit": None, "source": "Diabetes", "relations": [],
    }]}), encoding="utf-8")
    service = EavWorkflowService(object(), tmp_path / "runtime", base)
    filename, data = service.library_workbook()
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook["Criteria Library"]
    assert filename.startswith("criteriaLibrary_") and filename.endswith(".xlsx")
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.freeze_panes == "A2"
    assert sheet.max_row == 3
    assert {sheet.cell(2, 1).value, sheet.cell(3, 1).value} == {"immutable_base", "reviewed"}
    assert "Field Guide" in workbook.sheetnames


def test_validated_export_excludes_unready_rows(tmp_path: Path):
    base = tmp_path / "base.json"
    base.write_text(json.dumps({"immutable": True, "entries": []}), encoding="utf-8")
    snapshot = tmp_path / "criteriaLibrary_20200101_000000_000000.json"
    ready = {
        "criterion_id": "criterion-ready", "row_id": "criterion-ready",
        "parent_statement_id": "statement-1", "source_id": "statement-1",
        "trial_id": "NCT1", "protocol_id": "NCT1", "source_version": "v1",
        "source_retrieved_at": "2026-01-01T00:00:00+00:00", "protocol_hash": "sha256:test",
        "criterion_type": "inclusion", "domain": "person", "entity": "Age",
        "comparator": ">=", "value": "18", "unit": "years",
        "source": "Age >=18 years", "surrounding_context": "Age >=18 years",
        "human_review_status": "accepted", "reviewed_at": "2026-01-01T00:00:00+00:00",
        "extraction_confidence": 0.9,
    }
    pending = {**ready, "criterion_id": "criterion-pending", "row_id": "criterion-pending", "human_review_status": "pending", "reviewed_at": None}
    snapshot.write_text(json.dumps({"entries": [ready, pending]}), encoding="utf-8")
    service = EavWorkflowService(object(), tmp_path / "runtime", base)
    filename, data = service.validated_export()
    payload = json.loads(data)
    assert filename.startswith("validated_criteria_")
    assert payload["schema_version"] == "modern-autocrit.validated-criteria.v1"
    assert [row["criterion_id"] for row in payload["criteria"]] == ["criterion-ready"]
    assert payload["criteria"][0]["export_ready"] is True
    assert "metricspace_ready" not in payload["criteria"][0]


def test_review_rejects_inconsistent_boolean_group(tmp_path: Path):
    base = tmp_path / "base.json"
    base.write_text(json.dumps({"immutable": True, "entries": []}), encoding="utf-8")
    service = EavWorkflowService(object(), tmp_path / "runtime", base)
    source = "pregnant or breastfeeding"
    items = segment_criteria(f"Exclusion Criteria:\n- {source}", protocol_id="NCT1", source_version="v1")
    rows = normalize_rows({"rows": [
        _model_row(source, "pregnancy", "condition", logic="or"),
        _model_row(source, "breastfeeding", "observation", logic="or"),
    ]}, items, protocol_id="NCT1", source_version="v1", source_retrieved_at="2026-01-01T00:00:00+00:00")
    job = Job(
        job_id="job", trial_id="NCT1", input_mode="manual", status="completed",
        rows=rows, protocol_id="NCT1", source_version="v1",
        source_retrieved_at="2026-01-01T00:00:00+00:00", protocol_hash="sha256:test",
    )
    service._save(job)
    reviewed = [{**row, "review_status": "accepted"} for row in rows]
    reviewed[1]["logical_operator"] = "and"
    with pytest.raises(ValueError, match="consistent AND or OR"):
        service.apply_review("job", reviewed)


def test_reset_library_removes_snapshots_and_preserves_examples(tmp_path: Path):
    base = tmp_path / "base.json"
    base.write_text(json.dumps({"immutable": True, "entries": [{"trial_id": "EXAMPLE"}]}), encoding="utf-8")
    service = EavWorkflowService(object(), tmp_path / "runtime", base)
    config_snapshot = tmp_path / "criteriaLibrary_20200101_000000_000000.json"
    runtime_snapshot = service.libraries / "criteriaLibrary_20200101_000000_000000.csv"
    legacy_snapshot = tmp_path / "eavLibrary_20200101_000000_000000.json"
    unrelated = tmp_path / "keep.json"
    for path in (config_snapshot, runtime_snapshot, legacy_snapshot, unrelated):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("{}", encoding="utf-8")
    result = service.reset_library()
    assert result["removed"] == 3
    assert result["base_entries"] == 1
    assert base.is_file() and unrelated.is_file()
    assert not config_snapshot.exists() and not runtime_snapshot.exists() and not legacy_snapshot.exists()


def _model_row(source, entity, domain, *, comparator=None, value="present", logic="standalone"):
    return {
        "criterion_type": "inclusion", "domain": domain, "entity": entity,
        "comparator": comparator, "value": value, "unit": None, "negated": False,
        "temporal": None, "repetition": None, "qualifier": None,
        "logical_operator": logic, "source": source,
        "lower_bound": None, "lower_inclusive": None,
        "upper_bound": None, "upper_inclusive": None,
    }
