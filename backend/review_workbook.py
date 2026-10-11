"""Export and import human review workbooks for terminal cohort runs."""

from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from backend.api.app_state import create_app_state
from backend.batch_manifest import load_manifest, manifest_state_path
from backend.services.eav_workflow import EavWorkflowService

REVIEW_COLUMNS = [
    # Evidence stays visible while the reviewer moves across editable fields.
    "source", "parent_statement",
    # Human-checkable semantic fields.
    "criterion_type", "domain", "entity", "comparator", "value", "unit",
    "value_kind", "lower_bound", "lower_inclusive", "upper_bound",
    "upper_inclusive", "negated", "temporal", "repetition", "qualifier",
    "logical_operator", "review_status",
    # Automated quality signals.
    "validation_issues", "readiness_blockers", "extraction_confidence", "export_ready",
    # Stable identifiers and protocol provenance.
    "trial_id", "protocol_id", "source_uri", "source_version", "source_retrieved_at",
    "protocol_hash", "surrounding_context", "job_id", "criterion_id", "row_id",
    "source_id", "parent_statement_id", "parent_atom_count", "atom_index",
    "boolean_group_id",
]
BOOLEAN_COLUMNS = {"lower_inclusive", "upper_inclusive", "negated"}


def _service(root: Path) -> EavWorkflowService:
    return EavWorkflowService(
        create_app_state(), root / "runtime" / "criteria",
        root / "config" / "base_criterion_library.json",
    )


def _state(manifest: Path, root: Path) -> dict[str, Any]:
    state_path = manifest_state_path(manifest, root / "runtime" / "criteria")
    if not state_path.is_file():
        raise ValueError("No cohort state exists yet. Run extraction before exporting review.")
    return json.loads(state_path.read_text(encoding="utf-8"))


def export_review_workbook(manifest: Path, root: Path, output: Path | None = None) -> dict[str, Any]:
    """Write completed jobs from the selected manifest to an editable XLSX."""
    service = _service(root)
    state = _state(manifest, root)
    selected = set(load_manifest(manifest))
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Validation Instructions"
    sheet = workbook.create_sheet("Human Review")
    sheet.append(REVIEW_COLUMNS)
    completed_jobs = 0
    rows_written = 0
    trial_metrics: list[dict[str, Any]] = []
    issue_counts: Counter[str] = Counter()
    error_rows = warning_rows = 0
    inclusion_rows = exclusion_rows = 0
    confidences: list[float] = []
    for trial_id in load_manifest(manifest):
        record = state.get("jobs", {}).get(trial_id) or {}
        job_id = record.get("job_id")
        if not job_id:
            continue
        try:
            job = service.get(job_id)
        except KeyError:
            continue
        if job.get("status") != "completed" or job.get("trial_id") not in selected:
            continue
        completed_jobs += 1
        trial_row_count = trial_errors = trial_warnings = 0
        trial_confidences: list[float] = []
        for row in job.get("rows", []):
            values = dict(row)
            values["job_id"] = job_id
            values["trial_id"] = values.get("trial_id") or job.get("trial_id")
            values["protocol_hash"] = values.get("protocol_hash") or job.get("protocol_hash")
            values["review_status"] = "pending"
            issues = row.get("validation_issues", [])
            codes = [str(item.get("code") or item) for item in issues]
            values["validation_issues"] = " | ".join(
                codes
            )
            values["readiness_blockers"] = " | ".join(
                str(item) for item in row.get("readiness_blockers", [])
            )
            sheet.append([values.get(column) for column in REVIEW_COLUMNS])
            rows_written += 1
            trial_row_count += 1
            issue_counts.update(codes)
            has_error = any(item.get("severity") == "error" for item in issues if isinstance(item, dict))
            has_warning = any(item.get("severity") == "warning" for item in issues if isinstance(item, dict))
            error_rows += int(has_error); warning_rows += int(has_warning)
            trial_errors += int(has_error); trial_warnings += int(has_warning)
            inclusion_rows += int(row.get("criterion_type") == "inclusion")
            exclusion_rows += int(row.get("criterion_type") == "exclusion")
            confidence = row.get("extraction_confidence")
            if isinstance(confidence, (int, float)):
                confidences.append(float(confidence)); trial_confidences.append(float(confidence))
        trial_metrics.append({
            "trial_id": trial_id, "rows": trial_row_count,
            "error_rows": trial_errors, "warning_rows": trial_warnings,
            "mean_confidence": sum(trial_confidences) / len(trial_confidences) if trial_confidences else None,
        })

    header_fill = PatternFill("solid", fgColor="9F1724")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        "source": 70, "parent_statement": 70, "trial_id": 16,
        "criterion_type": 14, "domain": 15, "entity": 30,
        "comparator": 12, "value": 18, "unit": 14, "temporal": 26,
        "repetition": 24, "qualifier": 28, "logical_operator": 16,
        "surrounding_context": 70, "review_status": 17,
        "validation_issues": 28, "readiness_blockers": 28,
    }
    for index, column in enumerate(REVIEW_COLUMNS, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = widths.get(column, 18)
    status_column = REVIEW_COLUMNS.index("review_status") + 1
    status_validation = DataValidation(
        type="list", formula1='"pending,accepted,corrected"', allow_blank=False
    )
    sheet.add_data_validation(status_validation)
    if rows_written:
        status_validation.add(f"{sheet.cell(2, status_column).coordinate}:{sheet.cell(rows_written + 1, status_column).coordinate}")

    # The first sheet doubles as instructions and a read-only extraction-quality dashboard.
    instructions = [
        ("Workflow", "Review every row. Use accepted when unchanged, or corrected after editing one or more semantic fields."),
        ("Editable fields", "criterion_type, domain, entity, comparator, value, unit, bounds, negated, temporal, repetition, qualifier, and logical_operator."),
        ("Evidence", "Check source and surrounding_context before deciding. Do not edit identifiers or source/provenance columns."),
        ("Compound criteria", "Rows sharing parent_statement_id must retain complete atom_index ordering and one consistent AND/OR operator."),
        ("Apply", "Save the workbook, then run ./scripts/run_trial_cohort.py --apply-review path/to/workbook.xlsx."),
        ("Incomplete trials", "A trial is not applied until every one of its rows is marked accepted or corrected."),
    ]
    guide.append(["Topic", "Instructions"])
    for item in instructions:
        guide.append(item)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    guide.append([])
    guide.append(["Extraction quality overview", "Value"])
    overview = [
        ("Completed trials", completed_jobs),
        ("Extracted criteria", rows_written),
        ("Inclusion criteria", inclusion_rows),
        ("Exclusion criteria", exclusion_rows),
        ("Rows with blocking errors", error_rows),
        ("Rows with warnings", warning_rows),
        ("Mean extraction confidence", sum(confidences) / len(confidences) if confidences else None),
        ("Pending human review", rows_written),
        ("Initially export ready", 0),
    ]
    for item in overview:
        guide.append(item)
    guide.append([])
    guide.append(["Validation finding", "Count"])
    for code, count in sorted(issue_counts.items(), key=lambda item: (-item[1], item[0])):
        guide.append([code, count])
    guide.append([])
    guide.append(["Trial", "Criteria", "Rows with errors", "Rows with warnings", "Mean confidence"])
    for metric in trial_metrics:
        guide.append([
            metric["trial_id"], metric["rows"], metric["error_rows"],
            metric["warning_rows"], metric["mean_confidence"],
        ])
    for row_number in (1, len(instructions) + 3, len(instructions) + len(overview) + 5,
                       len(instructions) + len(overview) + len(issue_counts) + 7):
        for cell in guide[row_number]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
    guide.column_dimensions["C"].width = 20
    guide.column_dimensions["D"].width = 20
    guide.column_dimensions["E"].width = 20
    guide.freeze_panes = "A2"

    source_uri_column = REVIEW_COLUMNS.index("source_uri") + 1
    for row_number in range(2, rows_written + 2):
        cell = sheet.cell(row_number, source_uri_column)
        if cell.value:
            cell.hyperlink = str(cell.value)
            cell.style = "Hyperlink"

    if output is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output = root / f"{manifest.stem}_human_review_{stamp}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return {"path": str(output), "completed_jobs": completed_jobs, "rows": rows_written}


def _bool_or_none(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    raise ValueError(f"Expected a Boolean value, received {value!r}.")


def apply_review_workbook(manifest: Path, workbook_path: Path, root: Path) -> dict[str, Any]:
    """Validate and apply every fully reviewed trial in a workbook."""
    service = _service(root)
    state = _state(manifest, root)
    allowed_trials = set(load_manifest(manifest))
    workbook = load_workbook(workbook_path, data_only=True)
    if "Human Review" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a Human Review sheet.")
    sheet = workbook["Human Review"]
    headers = [cell.value for cell in sheet[1]]
    missing = [column for column in REVIEW_COLUMNS if column not in headers]
    if missing:
        raise ValueError("Review workbook is missing columns: " + ", ".join(missing))
    by_job: dict[str, list[dict[str, Any]]] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if not row.get("job_id"):
            continue
        if str(row.get("trial_id") or "").upper() not in allowed_trials:
            raise ValueError(f"Trial {row.get('trial_id')} is not in the selected manifest.")
        for column in BOOLEAN_COLUMNS:
            row[column] = _bool_or_none(row.get(column))
        by_job.setdefault(str(row["job_id"]), []).append(row)

    applied_jobs = 0
    applied_rows = 0
    incomplete: list[str] = []
    for job_id, rows in by_job.items():
        job = service.get(job_id)
        expected = len(job.get("rows", []))
        statuses = {str(row.get("review_status") or "").strip().casefold() for row in rows}
        if len(rows) != expected or not statuses <= {"accepted", "corrected"} or "pending" in statuses or "" in statuses:
            incomplete.append(str(job.get("trial_id") or job_id))
            continue
        service.apply_review(job_id, rows)
        applied_jobs += 1
        applied_rows += len(rows)
    return {"applied_jobs": applied_jobs, "applied_rows": applied_rows, "incomplete_trials": incomplete}
