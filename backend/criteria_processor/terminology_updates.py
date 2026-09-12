"""Audited working and verified EAV attribute-library updates."""

from __future__ import annotations

import csv
import json
import shutil
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from backend.criteria_processor.eav_schema import ALLOWED_ENTITIES, is_allowed_entity
from backend.criteria_processor.mapping_suggestions import audit_records_from_proposal
from backend.criteria_processor.models import MappingProposal, TerminologyAuditRecord
from backend.criteria_processor.text import clean_attribute_key


AUDIT_COLUMNS = list(TerminologyAuditRecord.model_fields)


def write_audit_csv(
    records: Iterable[TerminologyAuditRecord],
    path: str | Path,
) -> Path:
    """Write compact attribute-library audit records as editable CSV."""

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=AUDIT_COLUMNS)
        writer.writeheader()
        for record in records:
            row = record.model_dump(mode="json")
            for key, value in list(row.items()):
                if isinstance(value, list):
                    row[key] = json.dumps(value, ensure_ascii=False)
            writer.writerow(row)
    return output_path


def read_audit_csv(path: str | Path) -> list[TerminologyAuditRecord]:
    """Read reviewer-edited attribute-library audit records."""

    rows: list[TerminologyAuditRecord] = []
    with Path(path).open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            payload: dict[str, Any] = dict(row)
            for key in [
                "aliases",
                "example_trial_ids",
                "example_source_texts",
            ]:
                value = payload.get(key)
                if isinstance(value, str) and value.strip():
                    payload[key] = json.loads(value)
                else:
                    payload[key] = []
            if isinstance(payload.get("active"), str):
                payload["active"] = payload["active"].casefold() in {
                    "1",
                    "true",
                    "yes",
                    "y",
                    "active",
                }
            if payload.get("confidence") in {"", None}:
                payload["confidence"] = 0.0
            if payload.get("evidence_count") in {"", None}:
                payload["evidence_count"] = 1
            rows.append(TerminologyAuditRecord.model_validate(payload))
    return rows


def audit_records_from_proposals(
    proposals: Iterable[MappingProposal],
) -> list[TerminologyAuditRecord]:
    records: list[TerminologyAuditRecord] = []
    for proposal in proposals:
        records.extend(audit_records_from_proposal(proposal))
    return records


def create_working_library(
    *,
    verified_library_path: str | Path,
    working_library_path: str | Path,
    audit_file_path: str | Path,
    proposals: Iterable[MappingProposal],
) -> tuple[Path, Path]:
    """Create a working attribute library and its proposal audit file."""

    records = audit_records_from_proposals(proposals)
    write_audit_csv(records, audit_file_path)
    apply_audit_records_to_workbook(
        source_workbook=verified_library_path,
        output_workbook=working_library_path,
        records=records,
        allowed_review_statuses={"proposed", "approved", "revised"},
        metadata_status="working",
    )
    return Path(working_library_path), Path(audit_file_path)


def create_verified_library(
    *,
    current_verified_library_path: str | Path,
    new_verified_library_path: str | Path,
    audit_file_path: str | Path,
) -> Path:
    """Create a verified library from approved/revised audit rows only."""

    records = read_audit_csv(audit_file_path)
    apply_audit_records_to_workbook(
        source_workbook=current_verified_library_path,
        output_workbook=new_verified_library_path,
        records=records,
        allowed_review_statuses={"approved", "revised"},
        metadata_status="verified",
    )
    return Path(new_verified_library_path)


def apply_audit_records_to_workbook(
    *,
    source_workbook: str | Path,
    output_workbook: str | Path,
    records: Iterable[TerminologyAuditRecord],
    allowed_review_statuses: set[str],
    metadata_status: str,
) -> Path:
    """Copy an attribute workbook and apply selected audit changes."""

    source = Path(source_workbook)
    output = Path(output_workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    if source.resolve() != output.resolve():
        shutil.copy2(source, output)
    workbook = load_workbook(output)
    selected = [
        record
        for record in records
        if record.review_status in allowed_review_statuses
        and record.action != "no_change"
    ]
    for record in selected:
        if record.action != "add":
            raise NotImplementedError(
                "Only add actions are implemented for attribute-library "
                "updates."
            )
        if record.target_object == "attribute":
            _append_attribute(
                workbook["Attributes"],
                workbook["Aliases"],
                record,
            )
    for record in selected:
        if record.target_object == "value":
            _append_value(
                workbook["Values"],
                workbook["Aliases"],
                record,
            )
        elif record.target_object == "attribute_alias":
            _append_aliases(
                workbook["Aliases"],
                record,
                alias_type="attribute",
            )
        elif record.target_object == "value_alias":
            _append_aliases(
                workbook["Aliases"],
                record,
                alias_type="value",
            )
    _set_metadata_status(workbook["Metadata"], metadata_status)
    workbook.save(output)
    return output


def _headers(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value is not None
    }


def _rows_as_dicts(sheet: Worksheet) -> list[dict[str, Any]]:
    headers = _headers(sheet)
    rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(
            {
                name: row[index - 1]
                for name, index in headers.items()
                if index - 1 < len(row)
            }
        )
    return rows


def _append_by_headers(sheet: Worksheet, values: dict[str, Any]) -> None:
    headers = _headers(sheet)
    sheet.append([values.get(name, "") for name in headers])
    for table in sheet.tables.values():
        min_col, min_row, max_col, _max_row = range_boundaries(table.ref)
        table.ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{sheet.max_row}"
        )


def _append_attribute(
    sheet: Worksheet,
    aliases_sheet: Worksheet,
    record: TerminologyAuditRecord,
) -> None:
    attribute_id = record.attribute_id or record.object_id
    if not is_allowed_entity(record.parent_entity):
        raise ValueError(
            f"Attribute '{attribute_id}' must use one of the controlled "
            f"parent entities: {list(ALLOWED_ENTITIES)}"
        )
    existing_rows = _rows_as_dicts(sheet)
    existing = next(
        (
            row
            for row in existing_rows
            if str(row.get("attribute_id") or "") == attribute_id
        ),
        None,
    )
    if existing is not None:
        existing_parent = str(existing.get("parent_entity") or "").strip()
        proposed_parent = record.parent_entity.strip()
        if proposed_parent and existing_parent != proposed_parent:
            raise ValueError(
                f"Attribute '{attribute_id}' already belongs to fixed parent "
                f"entity '{existing_parent}', not '{proposed_parent}'."
            )
    else:
        _append_by_headers(
            sheet,
            {
                "attribute_id": attribute_id,
                "canonical_name": record.canonical_name,
                "parent_entity": record.parent_entity,
                "value_schema": record.value_schema,
                "canonical_unit": record.canonical_unit,
                "description": record.description or record.rationale,
                "active": record.active,
                "external_system": "",
                "external_code": "",
            },
        )
    _append_aliases(aliases_sheet, record, alias_type="attribute")


def _append_value(
    sheet: Worksheet,
    aliases_sheet: Worksheet,
    record: TerminologyAuditRecord,
) -> None:
    value_id = record.value_id or record.object_id
    existing = {
        str(row.get("value_id") or "")
        for row in _rows_as_dicts(sheet)
    }
    if value_id not in existing:
        sort_order = 1 + sum(
            1
            for row in _rows_as_dicts(sheet)
            if str(row.get("attribute_id") or "") == record.attribute_id
        )
        _append_by_headers(
            sheet,
            {
                "value_id": value_id,
                "attribute_id": record.attribute_id,
                "canonical_value": record.canonical_value,
                "sort_order": sort_order,
                "description": record.description or record.rationale,
                "active": record.active,
                "external_code": "",
            },
        )
    _append_aliases(aliases_sheet, record, alias_type="value")


def _append_aliases(
    sheet: Worksheet,
    record: TerminologyAuditRecord,
    *,
    alias_type: str,
) -> None:
    for alias in record.aliases:
        alias_key = clean_attribute_key(alias)
        if not alias_key:
            continue
        duplicate = False
        for row in _rows_as_dicts(sheet):
            if (
                clean_attribute_key(row.get("alias_text")) == alias_key
                and clean_attribute_key(row.get("source_entity"))
                == clean_attribute_key(record.source_entity)
                and str(row.get("attribute_id") or "") == record.attribute_id
                and str(row.get("value_id") or "") == record.value_id
                and str(row.get("alias_type") or "") == alias_type
            ):
                duplicate = True
                break
        if duplicate:
            continue
        alias_id = (
            f"alias_{record.value_id or record.attribute_id}_"
            f"{clean_attribute_key(record.source_entity).replace(' ', '_')}_"
            f"{clean_attribute_key(alias).replace(' ', '_')}"
        )
        _append_by_headers(
            sheet,
            {
                "alias_id": alias_id,
                "alias_text": alias,
                "source_entity": record.source_entity,
                "attribute_id": record.attribute_id,
                "value_id": record.value_id,
                "alias_type": alias_type,
                "status": "approved",
                "source": record.source,
                "active": record.active,
                "notes": (
                    f"{record.notes} change_id={record.change_id}; "
                    f"proposal_id={record.proposal_id}"
                ).strip(),
            },
        )


def _set_metadata_status(sheet: Worksheet, status: str) -> None:
    headers = _headers(sheet)
    key_col = headers.get("key")
    value_col = headers.get("value")
    if key_col is None or value_col is None:
        return
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=key_col).value == "status":
            sheet.cell(row=row, column=value_col).value = status
            return
    next_row = sheet.max_row + 1
    sheet.cell(row=next_row, column=key_col).value = "status"
    sheet.cell(row=next_row, column=value_col).value = status
