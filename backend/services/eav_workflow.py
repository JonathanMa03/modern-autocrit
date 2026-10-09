"""Extraction jobs and append-only structured-criterion snapshots."""

from __future__ import annotations

import csv
import base64
import hashlib
import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock, Thread
from typing import Any
from uuid import uuid4

from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from jsonschema import validate as validate_json_schema

from backend.api.app_state import AppState
from backend.atomic_io import atomic_write_json
from backend.eav_pipeline import ROW_SCHEMA, deterministic_criterion, extraction_prompt_for_items, interval_notation, normalize_rows, segment_criteria, semantic_relations
from backend.criterion_validation import readiness, validate_criterion
from backend.services.ctgov_v2_service import eligibility_text_from_study, fetch_study
from backend.services.llm.factory import create_llm_provider
from backend.services.llm.json_adapter import StructuredJsonAdapter


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


@dataclass
class Job:
    job_id: str
    trial_id: str
    input_mode: str
    status: str = "queued"
    stage: str = "queued"
    message: str = "Waiting to start."
    progress: int = 0
    created_at: str = field(default_factory=now)
    updated_at: str = field(default_factory=now)
    error: str | None = None
    rows: list[dict[str, Any]] = field(default_factory=list)
    events: list[dict[str, str]] = field(default_factory=list)
    protocol_id: str = ""
    source_version: str = ""
    source_uri: str | None = None
    source_retrieved_at: str | None = None
    protocol_hash: str = ""
    protocol_text: str = ""
    source_items: list[dict[str, Any]] = field(default_factory=list)
    checkpoint_rows: list[dict[str, Any]] = field(default_factory=list)
    completed_batches: int = 0


class EavWorkflowService:
    def __init__(self, state: AppState, root: Path, base_library: Path) -> None:
        self.state = state
        self.root = root
        self.jobs = root / "jobs"
        self.libraries = root / "libraries"
        self.base_library = base_library
        self.config_libraries = base_library.parent
        self._lock = Lock()
        self.jobs.mkdir(parents=True, exist_ok=True)
        self.libraries.mkdir(parents=True, exist_ok=True)

    def submit(self, *, input_mode: str, trial_id: str, protocol_text: str, protocol_pdf: str = "") -> dict[str, Any]:
        if input_mode not in {"manual", "nct", "pdf"}:
            raise ValueError("Input mode must be manual, nct, or pdf.")
        if input_mode == "manual" and not protocol_text.strip():
            raise ValueError("Paste protocol eligibility criteria.")
        if input_mode == "nct" and not trial_id.strip():
            raise ValueError("Enter an NCT ID.")
        if input_mode == "pdf" and not protocol_pdf:
            raise ValueError("Choose a protocol PDF.")
        job = Job(job_id=uuid4().hex, trial_id=trial_id.strip() or "MANUAL", input_mode=input_mode)
        self._save(job)
        Thread(target=self._run, args=(job.job_id, protocol_text, protocol_pdf), daemon=True).start()
        return asdict(job)

    def get(self, job_id: str) -> dict[str, Any]:
        path = self.jobs / f"{job_id}.json"
        if not path.is_file():
            raise KeyError(job_id)
        return json.loads(path.read_text(encoding="utf-8"))

    def list_jobs(self, *, limit: int = 100) -> list[dict[str, Any]]:
        jobs = []
        for path in sorted(self.jobs.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)[:limit]:
            job = json.loads(path.read_text(encoding="utf-8"))
            jobs.append({key: job.get(key) for key in (
                "job_id", "trial_id", "protocol_id", "source_version", "input_mode",
                "status", "stage", "progress", "created_at", "updated_at", "error",
            )})
        return jobs

    def resume(self, job_id: str) -> dict[str, Any]:
        job = Job(**self.get(job_id))
        if job.status == "completed":
            return asdict(job)
        job.status, job.error = "queued", None
        self._event(job, "resuming", f"Resuming after {job.completed_batches} completed extraction batches.", job.progress)
        Thread(target=self._run, args=(job.job_id, job.protocol_text, ""), daemon=True).start()
        return asdict(job)

    def _save(self, job: Job) -> None:
        job.updated_at = now()
        atomic_write_json(self.jobs / f"{job.job_id}.json", asdict(job))

    def _event(self, job: Job, stage: str, message: str, progress: int) -> None:
        job.stage, job.message, job.progress = stage, message, progress
        job.events.append({"timestamp": now(), "stage": stage, "message": message})
        self._save(job)

    @staticmethod
    def _eligibility_from_pdf(encoded_pdf: str) -> str:
        try:
            raw = base64.b64decode(encoded_pdf, validate=True)
            pages = [(page.extract_text() or "") for page in PdfReader(io.BytesIO(raw)).pages]
        except Exception as exc:
            raise ValueError(f"Could not read protocol PDF: {exc}") from exc
        numbered_inclusion = re.compile(
            r"^\s*\d+(?:\.\d+)+\s+Inclusion\s+criteria\s*$", re.I | re.M
        )
        start_page = next(
            (index for index, page in enumerate(pages) if numbered_inclusion.search(page)),
            None,
        )
        if start_page is None:
            start_page = next(
                (
                    index
                    for index, page in enumerate(pages)
                    if re.search(r"^\s*Key inclusion criteria\s*$", page, re.I | re.M)
                    and "......" not in page
                ),
                None,
            )
        if start_page is None:
            raise ValueError("The PDF did not contain a recognizable Inclusion or Exclusion Criteria section.")
        text = "\n".join(pages[start_page : start_page + 8])
        start = re.search(r"^\s*(?:\d+(?:\.\d+)+\s+)?(?:Key\s+)?Inclusion criteria\s*$", text, re.I | re.M)
        if start:
            text = text[start.start() :]
        # Page furniture inside a long protocol must not become part of the
        # preceding wrapped criterion.
        text = "\n".join(
            line for line in text.splitlines()
            if not (
                "Protocol Date:" in line
                or line.lstrip().startswith("Trial ID:")
                or "VV-TMF-" in line
                or re.search(r"\|\s*of\s+Protocol\b", line, re.I)
                or re.fullmatch(r"\s*Page:?\s+\d+\s+of\s+\d+\s*", line, re.I)
                or re.fullmatch(r"\s*(?:Confidential|(?:Status:\s*)?Final)\s*", line, re.I)
            )
        )
        headings = list(re.finditer(r"^\s*\d+(?:\.\d+)+\s+[^\n]+$", text, re.M))
        eligibility_headings_seen = 0
        for heading in headings:
            title = heading.group(0).casefold()
            if "inclusion criteria" in title or "exclusion criteria" in title:
                eligibility_headings_seen += 1
                continue
            if eligibility_headings_seen >= 2:
                text = text[: heading.start()]
                break
        return text

    def _run(self, job_id: str, protocol_text: str, protocol_pdf: str) -> None:
        job = Job(**self.get(job_id))
        try:
            job.status = "running"
            self._event(job, "ingestion", "Loading protocol eligibility criteria.", 15)
            if job.protocol_text:
                protocol_text = job.protocol_text
            elif job.input_mode == "nct":
                study = fetch_study(job.trial_id)
                protocol_text = eligibility_text_from_study(study)
                protocol = study.get("protocolSection") or {}
                status = protocol.get("statusModule") or {}
                version = (status.get("lastUpdatePostDateStruct") or {}).get("date") or (status.get("studyFirstPostDateStruct") or {}).get("date") or "unknown-date"
                job.protocol_id = job.trial_id.upper()
                job.source_uri = f"https://clinicaltrials.gov/study/{job.protocol_id}"
                job.source_retrieved_at = now()
                job.source_version = f"ctgov-{version}"
            elif job.input_mode == "pdf":
                self._event(job, "pdf", "Reading the full protocol PDF and locating eligibility sections.", 10)
                protocol_text = self._eligibility_from_pdf(protocol_pdf)
                if job.trial_id == "MANUAL":
                    job.trial_id = "PDF"
            digest = hashlib.sha256(protocol_text.encode("utf-8")).hexdigest()
            job.protocol_hash = f"sha256:{digest}"
            if not job.protocol_id:
                label = job.trial_id if job.trial_id not in {"MANUAL", "PDF"} else job.input_mode.upper()
                job.protocol_id = f"{label}-{digest[:12]}"
            if not job.source_version:
                job.source_version = f"sha256-{digest[:16]}"
            if not job.source_retrieved_at:
                job.source_retrieved_at = now()
            job.protocol_text = protocol_text
            source_items = job.source_items or segment_criteria(
                protocol_text, protocol_id=job.protocol_id, source_version=job.source_version
            )
            job.source_items = source_items
            self._save(job)
            if not source_items:
                raise ValueError("No eligibility criteria could be segmented from the input.")
            self._event(job, "located", f"Located {len(source_items)} atomic eligibility criteria.", 30)
            self._event(job, "extraction", f"Extracting 0 of {len(source_items)} criteria.", 35)
            provider = StructuredJsonAdapter(
                create_llm_provider(self.state.settings.llm, cost_monitor=self.state.cost_monitor),
                model=self.state.settings.llm.model,
            )
            extracted_rows = list(job.checkpoint_rows)
            batches = [source_items[index:index + 20] for index in range(0, len(source_items), 20)]
            for batch_number, batch in enumerate(batches[job.completed_batches:], start=job.completed_batches + 1):
                last_error: Exception | None = None
                payload = None
                for attempt in range(1, 3):
                    try:
                        payload = provider.generate_json(
                            extraction_prompt_for_items(batch), output_schema=ROW_SCHEMA
                        )
                        break
                    except Exception as exc:
                        last_error = exc
                        self._event(job, "retry", f"Batch {batch_number} attempt {attempt} failed; retrying.", 45)
                if payload is None:
                    raise RuntimeError(f"Structured extraction failed in batch {batch_number}: {last_error}")
                extracted_rows.extend(payload.get("rows") or [])
                job.checkpoint_rows = extracted_rows
                job.completed_batches = batch_number
                completed = min(batch_number * 20, len(source_items))
                self._event(job, "extraction", f"Extracted {completed} of {len(source_items)} criteria.", 40 + round(30 * completed / len(source_items)))
            self._event(job, "normalization", "Separating domains, constraints, constructs, and Boolean semantics.", 75)
            job.rows = normalize_rows(
                {"rows": extracted_rows}, source_items,
                protocol_id=job.protocol_id, source_version=job.source_version,
                source_uri=job.source_uri, source_retrieved_at=job.source_retrieved_at,
            )
            job.status = "completed"
            self._event(job, "completed", f"Extracted {len(job.rows)} criteria.", 100)
        except Exception as exc:
            job.status = "failed"
            job.error = f"{type(exc).__name__}: {exc}"
            self._event(job, "failed", job.error, 100)

    @staticmethod
    def _upgrade_legacy_entry(row: dict[str, Any]) -> dict[str, Any]:
        """Read former EAV snapshots without rewriting the historical files."""
        if "domain" in row:
            upgraded = dict(row)
            for key in ("lower_bound", "lower_inclusive", "upper_bound", "upper_inclusive"):
                upgraded.setdefault(key, None)
            upgraded.setdefault("comparator", None)
            upgraded.setdefault("value", None)
            upgraded.setdefault("unit", None)
            upgraded.setdefault("value_kind", "interval" if upgraded["lower_bound"] is not None and upgraded["upper_bound"] is not None else ("numerical" if upgraded.get("comparator") else "categorical"))
            upgraded.setdefault("negated", False)
            upgraded.setdefault("temporal", None)
            upgraded.setdefault("repetition", None)
            upgraded.setdefault("qualifier", None)
            upgraded.setdefault("logical_operator", "standalone")
            upgraded.setdefault("interval", interval_notation(
                upgraded["lower_bound"], upgraded["lower_inclusive"],
                upgraded["upper_bound"], upgraded["upper_inclusive"],
            ))
            upgraded.setdefault("relations", semantic_relations(
                comparator=upgraded.get("comparator") or ("interval" if upgraded.get("interval") else None),
                negated=bool(upgraded.get("negated")), temporal=upgraded.get("temporal"),
                repetition=upgraded.get("repetition"), qualifier=upgraded.get("qualifier"),
                logical_operator=str(upgraded.get("logical_operator") or "standalone"),
            ))
            legacy_material = json.dumps(upgraded, sort_keys=True, ensure_ascii=False, default=str)
            upgraded.setdefault("criterion_id", upgraded.get("row_id") or f"legacy-{hashlib.sha256(legacy_material.encode('utf-8')).hexdigest()[:16]}")
            upgraded.setdefault("parent_statement_id", upgraded.get("source_id") or upgraded["criterion_id"])
            upgraded.setdefault("parent_statement", upgraded.get("source") or "")
            upgraded.setdefault("parent_atom_count", 1)
            upgraded.setdefault("atom_index", 1)
            upgraded.setdefault("boolean_group_id", upgraded["parent_statement_id"])
            upgraded.setdefault("protocol_id", upgraded.get("trial_id") or "LEGACY")
            upgraded.setdefault("source_version", "legacy-unversioned")
            upgraded.setdefault("source_uri", None)
            upgraded.setdefault("source_retrieved_at", None)
            upgraded.setdefault("surrounding_context", upgraded.get("source") or "")
            upgraded.setdefault("extraction_confidence", 0.5)
            upgraded.setdefault("human_review_status", "accepted" if upgraded.get("reviewed_at") else "pending")
            upgraded["validation_issues"] = validate_criterion(upgraded)
            upgraded["export_ready"], upgraded["readiness_blockers"] = readiness(upgraded)
            upgraded.pop("metricspace_ready", None)
            upgraded.pop("metricspace_blockers", None)
            return upgraded
        semantic = deterministic_criterion(
            source=str(row.get("source") or ""),
            criterion_type=str(row.get("criterion_type") or "inclusion"),
            entity=str(row.get("entity") or ""),
            comparator=row.get("attribute"), value=row.get("value"),
        )
        return {
            **{key: value for key, value in row.items() if key != "attribute"},
            **semantic,
            "source_id": row.get("source_id") or row.get("row_id"),
        }

    def library(self) -> dict[str, Any]:
        base = json.loads(self.base_library.read_text(encoding="utf-8"))
        snapshots = sorted(self.config_libraries.glob("criteriaLibrary_*.json"))
        legacy = sorted(self.config_libraries.glob("eavLibrary_*.json"))
        selected = snapshots[-1] if snapshots else (legacy[-1] if legacy else None)
        latest = json.loads(selected.read_text(encoding="utf-8")) if selected else {"entries": []}
        base["entries"] = [self._upgrade_legacy_entry(row) for row in base.get("entries", [])]
        entries = [self._upgrade_legacy_entry(row) for row in latest.get("entries", [])]
        return {"base": base, "entries": entries, "snapshot": selected.name if selected else None}

    def library_workbook(self) -> tuple[str, bytes]:
        """Return an analysis-ready Excel export of the current library."""
        library = self.library()
        columns = [
            "library_source", "trial_id", "protocol_id", "source_version",
            "criterion_id", "parent_statement_id", "source_id", "row_id",
            "criterion_type", "domain", "entity", "comparator", "value",
            "unit", "value_kind", "lower_bound", "lower_inclusive",
            "upper_bound", "upper_inclusive", "interval", "negated",
            "temporal", "repetition", "qualifier", "logical_operator",
            "relations", "source", "surrounding_context", "extraction_confidence",
            "human_review_status", "validation_issues", "export_ready",
            "readiness_blockers", "reviewed_at",
        ]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Criteria Library"
        sheet.append(columns)
        rows = [
            ("immutable_base", row)
            for row in library["base"].get("entries", [])
        ] + [("reviewed", row) for row in library.get("entries", [])]
        for origin, row in rows:
            values = {**row, "library_source": origin}
            values["relations"] = " | ".join(str(item) for item in row.get("relations", []))
            values["validation_issues"] = " | ".join(str(item.get("code")) for item in row.get("validation_issues", []))
            values["export_ready"] = bool(row.get("export_ready"))
            values["readiness_blockers"] = " | ".join(str(item) for item in row.get("readiness_blockers", []))
            sheet.append([values.get(column) for column in columns])

        header_fill = PatternFill("solid", fgColor="9F1724")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths = {
            "library_source": 18, "trial_id": 16, "source_id": 16,
            "row_id": 34, "criterion_type": 14, "domain": 15,
            "entity": 30, "comparator": 12, "value": 18, "unit": 15,
            "value_kind": 14, "interval": 18, "temporal": 28,
            "repetition": 25, "qualifier": 30, "logical_operator": 16,
            "relations": 34, "source": 70, "reviewed_at": 27,
            "surrounding_context": 70, "validation_issues": 30,
            "readiness_blockers": 30, "criterion_id": 30,
            "parent_statement_id": 30, "source_version": 24,
        }
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 16)

        instructions = workbook.create_sheet("Field Guide")
        instructions.append(["Field", "How to interpret it"])
        guidance = [
            ("library_source", "immutable_base is packaged terminology; reviewed is human-reviewed trial output."),
            ("criterion_type", "Protocol inclusion or exclusion status; this does not reverse a comparator."),
            ("domain", "Chia-style clinical domain assigned to the entity."),
            ("entity", "Canonical clinical concept without threshold, unit, negation, timing, or frequency."),
            ("comparator / value / unit", "Literal quantitative constraint, or categorical present/absent value."),
            ("bounds / interval", "Normalized numerical endpoints used for overlap and restrictiveness comparisons."),
            ("negated", "True only when absence of the entity is required."),
            ("temporal / repetition / qualifier", "Clinically meaningful modifiers kept separate from the entity."),
            ("logical_operator", "How atomic rows from the same source combine: standalone, and, or."),
            ("source", "Verbatim evidence; use this when checking or manually labeling a row."),
            ("surrounding_context", "The parent statement plus adjacent criterion context retained from the protocol."),
            ("extraction_confidence", "Model-reported extraction confidence; it is not human approval or ground truth."),
            ("human_review_status", "pending, accepted unchanged, or corrected by a reviewer."),
            ("validation_issues", "Automated interval, negation, unit, temporal, repetition, and compound-scope findings."),
            ("export_ready", "True only after human review, complete provenance, and no blocking validation errors."),
            ("readiness_blockers", "Reasons a row is withheld from the validated downstream export."),
        ]
        for item in guidance:
            instructions.append(item)
        for cell in instructions[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        instructions.freeze_panes = "A2"
        instructions.column_dimensions["A"].width = 30
        instructions.column_dimensions["B"].width = 100
        for row in instructions.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        output = io.BytesIO()
        workbook.save(output)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"criteriaLibrary_{stamp}.xlsx", output.getvalue()

    def reset_library(self) -> dict[str, Any]:
        """Remove generated review snapshots while preserving the packaged base."""
        patterns = (
            "criteriaLibrary_*.json", "criteriaLibrary_*.csv",
            "eavLibrary_*.json", "eavLibrary_*.csv",
        )
        removed: list[str] = []
        with self._lock:
            for directory in (self.libraries, self.config_libraries):
                for pattern in patterns:
                    for path in directory.glob(pattern):
                        if path.is_file() and path.resolve() != self.base_library.resolve():
                            path.unlink()
                            removed.append(str(path))
        base = json.loads(self.base_library.read_text(encoding="utf-8"))
        return {
            "removed": len(removed),
            "base_entries": len(base.get("entries", [])),
            "message": "Reviewed criteria were removed; immutable examples were preserved.",
        }

    def validated_export(self) -> tuple[str, bytes]:
        """Create a neutral downstream export using only reviewed valid rows."""
        criteria = []
        for stored in self.library()["entries"]:
            if not stored.get("export_ready"):
                continue
            row = dict(stored)
            criteria.append(row)
        document = {
            "schema_version": "modern-autocrit.validated-criteria.v1",
            "created_at": now(),
            "criteria": criteria,
        }
        schema_path = Path(__file__).resolve().parents[1] / "schemas" / "validated_criteria_v1.schema.json"
        schema = json.loads(schema_path.read_text(encoding="utf-8"))
        validate_json_schema(instance=document, schema=schema)
        data = (json.dumps(document, indent=2, ensure_ascii=False) + "\n").encode("utf-8")
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        export_dir = self.root / "exports"
        export_dir.mkdir(parents=True, exist_ok=True)
        filename = f"validated_criteria_{stamp}.json"
        atomic_write_json(export_dir / filename, document)
        return filename, data

    def apply_review(self, job_id: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
        job = self.get(job_id)
        if job.get("status") != "completed":
            raise ValueError("Only completed jobs can be reviewed.")
        expected = {row["row_id"] for row in job.get("rows", [])}
        supplied = {str(row.get("row_id")) for row in rows}
        if expected != supplied:
            raise ValueError("Every extracted row must be accepted or corrected.")
        clean = []
        for row in rows:
            review_status = str(row.get("review_status") or row.get("human_review_status") or "")
            if review_status not in {"accepted", "corrected"}:
                raise ValueError("Every row must be accepted or corrected before apply.")
            criterion_type = str(row.get("criterion_type") or "").casefold()
            if criterion_type not in {"inclusion", "exclusion"}:
                raise ValueError("Criterion type must be inclusion or exclusion.")
            domain = str(row.get("domain") or "").casefold().strip()
            if domain not in {"observation", "condition", "person", "device", "drug", "visit", "procedure", "measurement"}:
                raise ValueError("Every reviewed row requires a valid Chia domain.")
            entity = str(row.get("entity") or "").strip()
            value = str(row.get("value") if row.get("value") is not None else "").strip()
            def optional_number(field: str) -> int | float | None:
                raw = row.get(field)
                if raw in {None, ""}: return None
                number = float(raw)
                return int(number) if number.is_integer() else number
            lower_bound = optional_number("lower_bound")
            upper_bound = optional_number("upper_bound")
            lower_inclusive = bool(row.get("lower_inclusive")) if lower_bound is not None else None
            upper_inclusive = bool(row.get("upper_inclusive")) if upper_bound is not None else None
            if not entity or (value == "" and lower_bound is None and upper_bound is None):
                raise ValueError("Entity and a value or interval bound are required for every reviewed row.")
            logical_operator = str(row.get("logical_operator") or "standalone").casefold()
            if logical_operator not in {"standalone", "and", "or"}:
                raise ValueError("Logical operator must be standalone, and, or or.")
            comparator = str(row.get("comparator") or "").strip() or None
            if comparator not in {None, ">", ">=", "<", "<=", "="}:
                raise ValueError("Comparator must be one of >, >=, <, <=, =, or blank.")
            temporal = str(row.get("temporal") or "").strip() or None
            repetition = str(row.get("repetition") or "").strip() or None
            qualifier = str(row.get("qualifier") or "").strip() or None
            negated = bool(row.get("negated"))
            reviewed = {
                "criterion_id": str(row.get("criterion_id") or row["row_id"]),
                "row_id": str(row["row_id"]), "trial_id": job["trial_id"],
                "source_id": str(row.get("source_id") or row["row_id"]),
                "parent_statement_id": str(row.get("parent_statement_id") or row.get("source_id") or row["row_id"]),
                "parent_statement": str(row.get("parent_statement") or row.get("source") or ""),
                "parent_atom_count": int(row.get("parent_atom_count") or 1),
                "atom_index": int(row.get("atom_index") or 1),
                "boolean_group_id": str(row.get("boolean_group_id") or row.get("parent_statement_id") or row.get("source_id") or row["row_id"]),
                "protocol_id": job.get("protocol_id") or job["trial_id"],
                "source_version": job.get("source_version") or "legacy-unversioned",
                "source_uri": job.get("source_uri"),
                "source_retrieved_at": job.get("source_retrieved_at"),
                "protocol_hash": job.get("protocol_hash"),
                "criterion_type": criterion_type, "domain": domain, "entity": entity,
                "comparator": comparator,
                "value": value, "unit": str(row.get("unit") or "").strip() or None,
                "value_kind": "interval" if lower_bound is not None and upper_bound is not None else str(row.get("value_kind") or "categorical"),
                "lower_bound": lower_bound, "lower_inclusive": lower_inclusive,
                "upper_bound": upper_bound, "upper_inclusive": upper_inclusive,
                "interval": interval_notation(lower_bound, lower_inclusive, upper_bound, upper_inclusive),
                "negated": negated,
                "temporal": temporal,
                "repetition": repetition,
                "qualifier": qualifier,
                "logical_operator": logical_operator,
                "relations": semantic_relations(comparator=comparator or ("interval" if lower_bound is not None or upper_bound is not None else None), negated=negated, temporal=temporal, repetition=repetition, qualifier=qualifier, logical_operator=logical_operator),
                "source": str(row.get("source") or "").strip(),
                "surrounding_context": str(row.get("surrounding_context") or row.get("source") or "").strip(),
                "extraction_confidence": float(row.get("extraction_confidence", 0.5)),
                "human_review_status": review_status,
                "reviewed_at": now(),
            }
            reviewed["validation_issues"] = validate_criterion(reviewed)
            reviewed["export_ready"], reviewed["readiness_blockers"] = readiness(reviewed)
            if not reviewed["export_ready"]:
                raise ValueError(
                    f"Criterion {reviewed['criterion_id']} is not ready for validated export: "
                    + ", ".join(reviewed["readiness_blockers"])
                )
            clean.append(reviewed)
        groups: dict[str, list[dict[str, Any]]] = {}
        for row in clean:
            groups.setdefault(row["parent_statement_id"], []).append(row)
        for parent_id, members in groups.items():
            expected_count = len(members)
            declared_counts = {row["parent_atom_count"] for row in members}
            indexes = {row["atom_index"] for row in members}
            operators = {row["logical_operator"] for row in members}
            if declared_counts != {expected_count} or indexes != set(range(1, expected_count + 1)):
                raise ValueError(f"Parent statement {parent_id} has inconsistent atom membership or ordering.")
            if expected_count > 1 and (len(operators) != 1 or operators <= {"standalone"}):
                raise ValueError(f"Parent statement {parent_id} must use one consistent AND or OR operator.")
        with self._lock:
            current = self.library()["entries"]
            by_key = {(r["trial_id"], r["source"], r["criterion_type"]): r for r in current}
            for row in clean:
                by_key[(row["trial_id"], row["source"], row["criterion_type"])] = row
            entries = sorted(by_key.values(), key=lambda r: (r["trial_id"], r["criterion_type"], r["entity"]))
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            json_path = self.libraries / f"criteriaLibrary_{stamp}.json"
            csv_path = self.libraries / f"criteriaLibrary_{stamp}.csv"
            config_json_path = self.config_libraries / json_path.name
            config_csv_path = self.config_libraries / csv_path.name
            document = {"schema_version": "modern-autocrit.structured-criteria.v4", "created_at": now(), "entries": entries}
            atomic_write_json(json_path, document)
            atomic_write_json(config_json_path, document)
            for target in (csv_path, config_csv_path):
                target.parent.mkdir(parents=True, exist_ok=True)
                with target.open("w", encoding="utf-8", newline="") as handle:
                    writer = csv.DictWriter(handle, fieldnames=["criterion_id", "row_id", "parent_statement_id", "parent_statement", "parent_atom_count", "atom_index", "boolean_group_id", "source_id", "trial_id", "protocol_id", "source_version", "source_uri", "source_retrieved_at", "protocol_hash", "criterion_type", "domain", "entity", "comparator", "value", "unit", "value_kind", "lower_bound", "lower_inclusive", "upper_bound", "upper_inclusive", "interval", "negated", "temporal", "repetition", "qualifier", "logical_operator", "relations", "source", "surrounding_context", "extraction_confidence", "human_review_status", "validation_issues", "export_ready", "readiness_blockers", "reviewed_at"], extrasaction="ignore")
                    writer.writeheader(); writer.writerows(entries)
        return {"added": len(clean), "total": len(entries), "json": json_path.name, "csv": csv_path.name, "config_json": str(config_json_path), "config_csv": str(config_csv_path), "entries": entries}
